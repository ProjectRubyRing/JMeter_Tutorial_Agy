# -*- coding: utf-8 -*-
"""
Sheet 04: Linux実行基盤構築 (Linux Setup & Transfer)
"""

from openpyxl.utils import get_column_letter

def build_sheet_linux_setup(wb, styles):
    ws = wb.create_sheet(title="04_Linux実行基盤構築")
    ws.views.sheetView[0].showGridLines = True

    styles["apply_title_banner"](
        ws,
        "04. Linux負荷テスト実行サーバの構築・JMX配置ガイド",
        "ディレクトリ設計・Java/JMeter配備・OSチューニング・Windowsからの転送手順",
        10
    )

    cur_r = 4
    styles["apply_section_header"](ws, cur_r, "1. Linux実行サーバの推奨ディレクトリ構造と初期セットアップ", 10)
    cur_r += 1

    linux_dirs = [
        ("/opt/jmeter/apache-jmeter-5.6.3", "JMeter本体プログラムの配置ディレクトリ (バイナリ、ライブラリ、binスクリプト群)"),
        ("/data/jmeter/scenarios", "Windows端末から転送されたJMXシナリオファイルの格納ディレクトリ"),
        ("/data/jmeter/csv", "シナリオで使用するパラメータCSVデータファイル (users.csv等) の格納ディレクトリ"),
        ("/data/jmeter/results", "テスト実行によって出力されるJTLサンプリング結果ファイル (.jtl) およびHTMLダッシュボード"),
        ("/data/jmeter/logs", "JMeter実行ログ (jmeter.log) およびCLI実行スクリプトの標準出力/標準エラーログ"),
        ("/data/jmeter/scripts", "テスト実行シェルスクリプト (run_load_test.sh, stop.sh等) の格納ディレクトリ")
    ]

    tbl_headers4 = ["ディレクトリパス", "用途・格納対象ファイル", "作成コマンド例"]
    ws.row_dimensions[cur_r].height = 24
    cols_map4 = [(1, 4), (5, 8), (9, 10)]
    for i, (sc, ec) in enumerate(cols_map4):
        if sc != ec:
            ws.merge_cells(start_row=cur_r, start_column=sc, end_row=cur_r, end_column=ec)
        for c in range(sc, ec + 1):
            cell = ws.cell(row=cur_r, column=c)
            cell.fill = styles["fill_tbl_header"]
            cell.border = styles["header_border"]
            if c == sc:
                cell.value = tbl_headers4[i]
                cell.font = styles["font_tbl_header"]
                cell.alignment = styles["align_center"]

    cur_r += 1
    for idx, d in enumerate(linux_dirs):
        ws.row_dimensions[cur_r].height = 24
        fill_curr = styles["fill_zebra"] if idx % 2 == 1 else styles["fill_white"]

        # Path
        ws.merge_cells(start_row=cur_r, start_column=1, end_row=cur_r, end_column=4)
        for c in range(1, 5):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = fill_curr
            c_cell.border = styles["cell_border"]
            if c == 1:
                c_cell.value = d[0]
                c_cell.font = styles["font_code"]
                c_cell.alignment = styles["align_left"]

        # Purpose
        ws.merge_cells(start_row=cur_r, start_column=5, end_row=cur_r, end_column=8)
        for c in range(5, 9):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = fill_curr
            c_cell.border = styles["cell_border"]
            if c == 5:
                c_cell.value = d[1]
                c_cell.font = styles["font_body"]
                c_cell.alignment = styles["align_left"]

        # Command
        ws.merge_cells(start_row=cur_r, start_column=9, end_row=cur_r, end_column=10)
        for c in range(9, 11):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = fill_curr
            c_cell.border = styles["cell_border"]
            if c == 9:
                c_cell.value = f"sudo mkdir -p {d[0]}"
                c_cell.font = styles["font_code"]
                c_cell.alignment = styles["align_left"]
        cur_r += 1

    # Section 2: Linux Setup Steps
    cur_r += 1
    styles["apply_section_header"](ws, cur_r, "2. Linux実行サーバの構築コマンド（JDK/JMeter/OSパラメータ）", 10)
    cur_r += 1

    linux_setup = [
        ("1", "OpenJDK 17 のインストール", 
         "# RHEL / Rocky Linux の場合:\n"
         "sudo dnf install -y java-17-openjdk java-17-openjdk-devel\n"
         "# Ubuntu / Debian の場合:\n"
         "sudo apt update && sudo apt install -y openjdk-17-jdk\n"
         "# バージョン確認:\n"
         "java -version", "Java 17 OpenJDK Runtime Environment が表示されること"),
        ("2", "JMeter のダウンロードと解凍・配置",
         "cd /tmp\n"
         "wget https://archive.apache.org/dist/jmeter/binaries/apache-jmeter-5.6.3.tgz\n"
         "sudo mkdir -p /opt/jmeter\n"
         "sudo tar -zxvf apache-jmeter-5.6.3.tgz -C /opt/jmeter/\n"
         "sudo chown -R $USER:$USER /opt/jmeter\n"
         "sudo chmod +x /opt/jmeter/apache-jmeter-5.6.3/bin/*", "jmeter実行スクリプトに実行権限が付与されていること"),
        ("3", "setenv.sh の作成（JVMヒープとGCの最適化）",
         "cat << 'EOF' > /opt/jmeter/apache-jmeter-5.6.3/bin/setenv.sh\n"
         "#!/bin/bash\n"
         "# サーバ実メモリ8GB想定: ヒープに4GB、G1GCを指定\n"
         "export JVM_ARGS=\"-Xms4g -Xmx4g -XX:MaxMetaspaceSize=512m -XX:+UseG1GC -XX:MaxGCPauseMillis=100 -XX:G1ReservePercent=15\"\n"
         "EOF\n"
         "chmod +x /opt/jmeter/apache-jmeter-5.6.3/bin/setenv.sh", "高負荷時もGC停止時間を最小化し、OOMを防止"),
        ("4", "環境変数 (PATH / JMETER_HOME) の登録",
         "cat << 'EOF' >> ~/.bashrc\n"
         "export JMETER_HOME=/opt/jmeter/apache-jmeter-5.6.3\n"
         "export PATH=$PATH:$JMETER_HOME/bin\n"
         "EOF\n"
         "source ~/.bashrc\n"
         "jmeter -v", "任意のディレクトリから jmeter -v でロゴが表示されること"),
        ("5", "OSカーネル・ファイルディスクリプタのチューニング",
         "# ファイルオープン数上限 (ulimit) の拡張:\n"
         "echo \"* soft nofile 65535\" | sudo tee -a /etc/security/limits.conf\n"
         "echo \"* hard nofile 65535\" | sudo tee -a /etc/security/limits.conf\n"
         "# エフェメラルポート拡張とTIME_WAITソケットの再利用:\n"
         "cat << 'EOF' | sudo tee -a /etc/sysctl.conf\n"
         "net.ipv4.ip_local_port_range = 1024 65535\n"
         "net.ipv4.tcp_tw_reuse = 1\n"
         "EOF\n"
         "sudo sysctl -p", "ulimit -n で 65535 が表示され、ソケット枯渇を防止できる状態")
    ]

    for step in linux_setup:
        ws.row_dimensions[cur_r].height = 54
        ws.cell(row=cur_r, column=1, value=step[0]).font = styles["font_body_bold"]
        ws.cell(row=cur_r, column=1).fill = styles["fill_card"]
        ws.cell(row=cur_r, column=1).border = styles["cell_border"]
        ws.cell(row=cur_r, column=1).alignment = styles["align_center"]

        ws.cell(row=cur_r, column=2, value=step[1]).font = styles["font_body_bold"]
        ws.cell(row=cur_r, column=2).fill = styles["fill_card"]
        ws.cell(row=cur_r, column=2).border = styles["cell_border"]
        ws.cell(row=cur_r, column=2).alignment = styles["align_left"]

        ws.merge_cells(start_row=cur_r, start_column=3, end_row=cur_r, end_column=7)
        for c in range(3, 8):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = styles["fill_code"]
            c_cell.border = styles["cell_border"]
            if c == 3:
                c_cell.value = step[2]
                c_cell.font = styles["font_code"]
                c_cell.alignment = styles["align_left"]

        ws.merge_cells(start_row=cur_r, start_column=8, end_row=cur_r, end_column=10)
        for c in range(8, 11):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = styles["fill_white"]
            c_cell.border = styles["cell_border"]
            if c == 8:
                c_cell.value = step[3]
                c_cell.font = styles["font_body"]
                c_cell.alignment = styles["align_left"]
        cur_r += 1

    # Section 3: Transfer Guide
    cur_r += 1
    styles["apply_section_header"](ws, cur_r, "3. Windows端末からLinux実行サーバへのファイル転送手順 (SCP / WinSCP)", 10)
    cur_r += 1

    transfer_methods = [
        ("PowerShell (SCPコマンド)", "CLI自動化", 
         "# 1. シナリオファイル (.jmx) の転送:\n"
         "scp -P 22 .\\scenarios\\sample_cognito_apigw.jmx user@linux-host:/data/jmeter/scenarios/\n"
         "# 2. パラメータCSVデータの転送:\n"
         "scp -P 22 .\\data\\sample_users.csv user@linux-host:/data/jmeter/csv/\n"
         "# 3. 実行スクリプトの転送:\n"
         "scp -P 22 .\\scripts\\run_load_test.sh user@linux-host:/data/jmeter/scripts/\n"
         "※SSHキー認証の場合は「-i C:\\path\\to\\id_rsa」を付与します。", "CI/CDや定期実行の自動化に適している。コマンド一発で転送完了。"),
        ("WinSCP (GUIツール)", "GUI手動操作",
         "1. WinSCPを起動し、ホスト名、ポート、ユーザー名、秘密鍵を設定して接続。\n"
         "2. 左側ペイン (Windowsローカル) で作成した .jmx および .csv を選択。\n"
         "3. 右側ペイン (Linuxリモート) で「/data/jmeter/scenarios/」または「/data/jmeter/csv/」を開く。\n"
         "4. ドラッグ＆ドロップでファイルを転送。\n"
         "※注意: 改行コードが混在する場合は、転送モードを「バイナリ」にして転送すること。", "初心者が視覚的にファイル配置を確認しながら安全に作業するのに最適。")
    ]

    for tm in transfer_methods:
        ws.row_dimensions[cur_r].height = 54
        ws.cell(row=cur_r, column=1, value=tm[0]).font = styles["font_body_bold"]
        ws.cell(row=cur_r, column=1).fill = styles["fill_card"]
        ws.cell(row=cur_r, column=1).border = styles["cell_border"]
        ws.cell(row=cur_r, column=1).alignment = styles["align_left"]

        ws.cell(row=cur_r, column=2, value=tm[1]).font = styles["font_body_bold"]
        ws.cell(row=cur_r, column=2).fill = styles["fill_card"]
        ws.cell(row=cur_r, column=2).border = styles["cell_border"]
        ws.cell(row=cur_r, column=2).alignment = styles["align_center"]

        ws.merge_cells(start_row=cur_r, start_column=3, end_row=cur_r, end_column=7)
        for c in range(3, 8):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = styles["fill_code"]
            c_cell.border = styles["cell_border"]
            if c == 3:
                c_cell.value = tm[2]
                c_cell.font = styles["font_code"]
                c_cell.alignment = styles["align_left"]

        ws.merge_cells(start_row=cur_r, start_column=8, end_row=cur_r, end_column=10)
        for c in range(8, 11):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = styles["fill_white"]
            c_cell.border = styles["cell_border"]
            if c == 8:
                c_cell.value = tm[3]
                c_cell.font = styles["font_body"]
                c_cell.alignment = styles["align_left"]
        cur_r += 1

    # Image Area for Sheet 4
    cur_r += 1
    styles["apply_section_header"](ws, cur_r, "4. サーバ基盤構築・ファイル配置証跡貼付エリア", 10)
    cur_r += 1
    styles["apply_image_placeholder"](
        ws, cur_r, 1, cur_r + 8, 5,
        "Windows から Linux への SCP / WinSCP 転送実行画面",
        "PowerShellでのSCP転送完了ログ、またはWinSCPでJMX/CSVファイルが配置された画面キャプチャ"
    )
    styles["apply_image_placeholder"](
        ws, cur_r, 6, cur_r + 8, 10,
        "Linux サーバ上での資材配置確認 (ls -la) & ulimit 実行結果",
        "/data/jmeter 配下のディレクトリ構成確認 (ls -la) および ulimit -n の 65535 表示確認証跡"
    )
    cur_r += 9

    # Set column widths
    col_widths4 = [12, 16, 14, 14, 16, 16, 16, 16, 16, 12]
    for idx, w in enumerate(col_widths4, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
