# -*- coding: utf-8 -*-
"""
Sheet 06: 結果取得とExcel取込 (Result Collection & Excel Import)
"""

from openpyxl.utils import get_column_letter

def build_sheet_excel_import(wb, styles):
    ws = wb.create_sheet(title="06_結果取得とExcel取込")
    ws.views.sheetView[0].showGridLines = True

    styles["apply_title_banner"](
        ws,
        "06. 負荷テスト結果回収・JTLファイル仕様・Excel取込手順",
        "LinuxからのSCP回収・JTL全17フィールド定義・Power Query取込手順・エポックミリ秒日時変換",
        10
    )

    cur_r = 4
    styles["apply_section_header"](ws, cur_r, "1. Linux実行サーバからのテスト結果ファイル回収手順", 10)
    cur_r += 1

    collect_cmd = (
        "【PowerShell でのテスト結果回収コマンド】\n"
        "# 1. JTLサンプリングログの回収:\n"
        "scp -P 22 user@linux-host:/data/jmeter/results/*.jtl C:\\PerformanceTest\\Results\\\n"
        "# 2. JMeterシステムログの回収:\n"
        "scp -P 22 user@linux-host:/data/jmeter/logs/*.log C:\\PerformanceTest\\Logs\\\n"
        "# 3. HTMLダッシュボードレポートの回収 (フォルダごと):\n"
        "scp -P 22 -r user@linux-host:/data/jmeter/results/report_* C:\\PerformanceTest\\Reports\\"
    )
    ws.merge_cells(start_row=cur_r, start_column=1, end_row=cur_r + 2, end_column=10)
    for r in range(cur_r, cur_r + 3):
        for c in range(1, 11):
            cell = ws.cell(row=r, column=c)
            cell.fill = styles["fill_code"]
            cell.border = styles["cell_border"]
            if r == cur_r and c == 1:
                cell.value = collect_cmd
                cell.font = styles["font_code"]
                cell.alignment = styles["align_left"]
    ws.row_dimensions[cur_r].height = 20
    ws.row_dimensions[cur_r + 1].height = 20
    ws.row_dimensions[cur_r + 2].height = 24
    cur_r += 4

    # Section 2: JTL File Field Definitions
    styles["apply_section_header"](ws, cur_r, "2. JMeter JTL (CSV形式) ファイルの全17フィールド完全仕様定義", 10)
    cur_r += 1

    jtl_fields = [
        ("timeStamp", "エポックミリ秒 (UNIX Time)", "リクエストが開始された時刻 (ミリ秒単位)。ExcelでJST日時に変換可能。"),
        ("elapsed", "経過時間 / 応答時間 (ms)", "リクエスト送信開始からレスポンス受信完了までの所要時間。主要分析対象。"),
        ("label", "サンプラー表示名", "シナリオ内で命名したHTTPサンプラーの名前 (例: 01_Auth_CognitoToken)。"),
        ("responseCode", "HTTPステータスコード", "サーバから返却されたHTTPコード (200, 400, 401, 403, 500, 504等)。"),
        ("responseMessage", "HTTPレスポンスメッセージ", "ステータスコードに付随するメッセージ (OK, Unauthorized, Forbidden等)。"),
        ("threadName", "実行スレッド識別名", "スレッドグループ名および仮想ユーザー番号 (例: TG_API 1-15)。"),
        ("dataType", "データ種別", "返却データのMIMEタイプ分類 (通常は text)。"),
        ("success", "成否フラグ (true / false)", "全アサーションに合格し、HTTPコードが正常であれば true、失敗なら false。"),
        ("failureMessage", "アサーション失敗理由", "アサーションが不合格となった場合の具体的なエラー内容。"),
        ("bytes", "受信バイト数", "レスポンス本文およびヘッダを含めてサーバから受信した全バイト数。"),
        ("sentBytes", "送信バイト数", "リクエストURL、ヘッダ、BodyデータとしてJMeterが送信した全バイト数。"),
        ("grpThreads", "自グループのアクティブスレッド数", "サンプリング瞬間に所属スレッドグループ内で稼働していたスレッド数。"),
        ("allThreads", "全体のアクティブスレッド数", "全スレッドグループを合わせたJMeterプロセス全体のアクティブスレッド数。"),
        ("URL", "リクエスト先完全URL", "実際にアクセスされたFQDNおよびパス、クエリパラメータ。"),
        ("Latency", "ネットワークレイテンシ (ms)", "リクエスト送信開始からレスポンスの「最初の1バイト」を受信するまでの時間。"),
        ("IdleTime", "アイドル時間 (ms)", "タイマーなどによる意図的な待機時間。通常は 0。"),
        ("Connect", "TCP/TLS接続確立時間 (ms)", "TCP 3-wayハンドシェイクおよびSSL/TLSハンドシェイクに要した時間。")
    ]

    tbl_headers6 = ["フィールド名 (CSVヘッダ)", "日本語項目名・単位", "フィールド定義・分析上の着眼点"]
    ws.row_dimensions[cur_r].height = 24
    cols_map6 = [(1, 3), (4, 5), (6, 10)]
    for i, (sc, ec) in enumerate(cols_map6):
        if sc != ec:
            ws.merge_cells(start_row=cur_r, start_column=sc, end_row=cur_r, end_column=ec)
        for c in range(sc, ec + 1):
            cell = ws.cell(row=cur_r, column=c)
            cell.fill = styles["fill_tbl_header"]
            cell.border = styles["header_border"]
            if c == sc:
                cell.value = tbl_headers6[i]
                cell.font = styles["font_tbl_header"]
                cell.alignment = styles["align_center"]

    cur_r += 1
    for idx, jf in enumerate(jtl_fields):
        ws.row_dimensions[cur_r].height = 20
        fill_curr = styles["fill_zebra"] if idx % 2 == 1 else styles["fill_white"]

        # Field
        ws.merge_cells(start_row=cur_r, start_column=1, end_row=cur_r, end_column=3)
        for c in range(1, 4):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = fill_curr
            c_cell.border = styles["cell_border"]
            if c == 1:
                c_cell.value = jf[0]
                c_cell.font = styles["font_code"]
                c_cell.alignment = styles["align_left"]

        # Japanese
        ws.merge_cells(start_row=cur_r, start_column=4, end_row=cur_r, end_column=5)
        for c in range(4, 6):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = fill_curr
            c_cell.border = styles["cell_border"]
            if c == 4:
                c_cell.value = jf[1]
                c_cell.font = styles["font_body_bold"]
                c_cell.alignment = styles["align_left"]

        # Description
        ws.merge_cells(start_row=cur_r, start_column=6, end_row=cur_r, end_column=10)
        for c in range(6, 11):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = fill_curr
            c_cell.border = styles["cell_border"]
            if c == 6:
                c_cell.value = jf[2]
                c_cell.font = styles["font_body"]
                c_cell.alignment = styles["align_left"]
        cur_r += 1

    # Section 3: Excel Import & Epoch Conversion
    cur_r += 1
    styles["apply_section_header"](ws, cur_r, "3. ExcelでのJTL取り込み手順（Power Query）とエポックミリ秒変換数式", 10)
    cur_r += 1

    pq_steps = [
        ("Step 1: Power Queryによる取込", 
         "1. Excel上部「データ」タブ →「データの取得」→「ファイルから」→「テキスト/CSVから」を選択。\n"
         "2. 回収した .jtl ファイルを選択し、「インポート」をクリック。\n"
         "3. プレビュー画面で以下を設定:\n"
         "   - 元のファイル (文字コード): 「65001: Unicode (UTF-8)」を選択。\n"
         "   - 区切り記号: 「コンマ」を選択。\n"
         "4. 右下の「データの変換」をクリックしてPower Queryエディターを開く。"),
        ("Step 2: 型の確認とシート読込",
         "1. elapsed, bytes, sentBytes, Latency, Connect 列のデータ型が「整数」になっていることを確認。\n"
         "2. success 列が「論理値 (TRUE/FALSE)」になっていることを確認。\n"
         "3. 「閉じて読み込む」をクリックし、Excelシート上にテーブルとして展開。"),
        ("Step 3: エポックミリ秒のJST変換",
         "JTLの timeStamp は 1970年1月1日からの通算ミリ秒 (例: 1728000000000) です。\n"
         "【Excel変換数式】(タイムスタンプが A2 セルの場合):\n"
         "  = ((A2 / 1000) / 86400) + DATE(1970, 1, 1) + (9 / 24)\n"
         "【解説】\n"
         "・A2 / 1000 : ミリ秒を秒に変換\n"
         "・/ 86400 : 秒を日 (1日=86400秒) に変換\n"
         "・+ DATE(1970,1,1) : UNIXエポック基準日を加算\n"
         "・+ (9/24) : UTCから日本標準時 (JST = UTC+9時間) へ時差補正\n"
         "※入力したセルの表示形式を「yyyy/mm/dd hh:mm:ss」に設定します。")
    ]

    for ps in pq_steps:
        ws.row_dimensions[cur_r].height = 54
        ws.merge_cells(start_row=cur_r, start_column=1, end_row=cur_r, end_column=3)
        for c in range(1, 4):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = styles["fill_card"]
            c_cell.border = styles["cell_border"]
            if c == 1:
                c_cell.value = ps[0]
                c_cell.font = styles["font_body_bold"]
                c_cell.alignment = styles["align_left"]

        ws.merge_cells(start_row=cur_r, start_column=4, end_row=cur_r, end_column=10)
        for c in range(4, 11):
            c_cell = ws.cell(row=cur_r, column=c)
            c_cell.fill = styles["fill_white"]
            c_cell.border = styles["cell_border"]
            if c == 4:
                c_cell.value = ps[1]
                c_cell.font = styles["font_body"]
                c_cell.alignment = styles["align_left"]
        cur_r += 1

    # Image Area for Sheet 6
    cur_r += 1
    styles["apply_section_header"](ws, cur_r, "4. 結果回収・Excel取込証跡貼付エリア", 10)
    cur_r += 1
    styles["apply_image_placeholder"](
        ws, cur_r, 1, cur_r + 8, 5,
        "Windows 端末での SCP 結果回収実行画面",
        "PowerShell上でLinuxサーバからJTLファイルおよびログファイルをローカル端末へ回収完了した画面キャプチャ"
    )
    styles["apply_image_placeholder"](
        ws, cur_r, 6, cur_r + 8, 10,
        "Excel Power Query による JTL (CSV) インポート画面",
        "Power QueryでUTF-8文字コード・カンマ区切りを指定し、JTLデータが正しくプレビュー・展開された証跡キャプチャ"
    )
    cur_r += 9

    # Set column widths
    col_widths6 = [12, 14, 14, 14, 16, 16, 16, 16, 16, 12]
    for idx, w in enumerate(col_widths6, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
