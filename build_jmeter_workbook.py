# -*- coding: utf-8 -*-
"""
Apache JMeter End-to-End Practical Guide and Performance Analysis Report Generator
Generates: JMeter_EndToEnd_Guide_And_Report.xlsx
Font: Meiryo UI everywhere
Color palette: Professional Monotone Base
Gridlines: Explicitly enabled
Image insertion placeholder areas with dashed borders and clear instructions.
All 9 Sheets:
  00_表紙・全体概要
  01_Win環境構築
  02_シナリオ作成手法
  03_Cognito・APIGW連携 (NEW!)
  04_Linux実行基盤構築
  05_CLI実行と監視
  06_結果取得とExcel取込
  07_性能分析レポート
  08_確認観点と注意点
"""

import os
import sys
import openpyxl

# Add scripts directory to sys.path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "scripts"))

from workbook.common_styles import get_workbook_styles
from workbook.sheet00_cover import build_sheet_cover
from workbook.sheet01_win_setup import build_sheet_win_setup
from workbook.sheet02_scenario_gui import build_sheet_scenario_gui
from workbook.sheet03_cognito_apigw import build_sheet_cognito_apigw
from workbook.sheet04_linux_setup import build_sheet_linux_setup
from workbook.sheet05_cli_execution import build_sheet_cli_execution
from workbook.sheet06_excel_import import build_sheet_excel_import
from workbook.sheet07_performance_report import build_sheet_performance_report
from workbook.sheet08_checklist import build_sheet_checklist

def build_workbook():
    wb = openpyxl.Workbook()
    styles = get_workbook_styles()

    print("Building Sheet 00: 00_表紙・全体概要...")
    build_sheet_cover(wb, styles)

    print("Building Sheet 01: 01_Win環境構築...")
    build_sheet_win_setup(wb, styles)

    print("Building Sheet 02: 02_シナリオ作成手法...")
    build_sheet_scenario_gui(wb, styles)

    print("Building Sheet 03: 03_Cognito・APIGW連携...")
    build_sheet_cognito_apigw(wb, styles)

    print("Building Sheet 04: 04_Linux実行基盤構築...")
    build_sheet_linux_setup(wb, styles)

    print("Building Sheet 05: 05_CLI実行と監視...")
    build_sheet_cli_execution(wb, styles)

    print("Building Sheet 06: 06_結果取得とExcel取込...")
    build_sheet_excel_import(wb, styles)

    print("Building Sheet 07: 07_性能分析レポート...")
    build_sheet_performance_report(wb, styles)

    print("Building Sheet 08: 08_確認観点と注意点...")
    build_sheet_checklist(wb, styles)

    # Remove default initial sheet
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    output_filename = "JMeter_EndToEnd_Guide_And_Report.xlsx"
    wb.save(output_filename)
    print(f"\n[SUCCESS] Workbook successfully generated and saved as: {output_filename}")
    print(f"Total Sheets: {len(wb.sheetnames)}")
    for i, sname in enumerate(wb.sheetnames):
        print(f"  [{i}] {sname}")

if __name__ == "__main__":
    build_workbook()
